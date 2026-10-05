from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
import json

root = Path(__file__).resolve().parents[1]
document = root / 'output/Assignment_3_Immutable_Book_Report.docx'
doc = Document(document)

def read(path):
    data = path.read_bytes()
    return data.decode('utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig')

sources = ['Book.cs', 'DefaultBook.cs', 'Program.cs', 'Checks.cs', 'verify.ps1',
           'IFP3.csproj', 'baseline/OriginalSnippet.txt', 'baseline/Program.cs',
           'baseline/Baseline.csproj']
for file in sources:
    start = next(i for i, p in enumerate(doc.paragraphs) if p.text == file)
    actual = []
    for paragraph in doc.paragraphs[start + 1:]:
        if paragraph.style.name != 'Code':
            break
        actual.append(paragraph.text.rstrip())
    expected = [line.expandtabs(4).rstrip() for line in read(root / file).rstrip().splitlines()]
    assert actual == expected, f'Code listing mismatch: {file}'

assert 'MISSING' not in '\n'.join(p.text for p in doc.paragraphs)
archive = root / 'output/Assignment_3_Source.zip'
files = sources + ['README.md'] + [str(p.relative_to(root)).replace('\\', '/')
                                  for p in sorted((root / 'evidence').glob('*')) if p.is_file()]
with ZipFile(archive, 'w', ZIP_DEFLATED) as z:
    for file in files:
        z.write(root / file, file)
with ZipFile(archive) as z:
    assert z.testzip() is None
    assert set(z.namelist()) == set(files)
    for file in files:
        assert z.read(file) == (root / file).read_bytes()

print(json.dumps({'complete_code_listings_verified': len(sources),
                  'archive_files': len(files), 'archive_bytes': archive.stat().st_size,
                  'docx_bytes': document.stat().st_size}, indent=2))
