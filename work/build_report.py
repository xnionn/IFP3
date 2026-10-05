from pathlib import Path
import json
import re
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'Assignment_3_Immutable_Book_Report.docx'
def read(path):
    data = path.read_bytes()
    return data.decode('utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.68)
sec.left_margin = sec.right_margin = Inches(.73)
sec.header_distance = sec.footer_distance = Inches(.3)
styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    styles[name].font.name = 'Calibri'
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Normal'].paragraph_format.line_spacing = 1.08
styles['Title'].font.size = Pt(24)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(12)
styles['Subtitle'].font.size = Pt(12)
styles['Heading 1'].font.size = Pt(16)
styles['Heading 1'].paragraph_format.space_before = Pt(12)
styles['Heading 1'].paragraph_format.space_after = Pt(8)
styles['Heading 2'].font.size = Pt(12)
styles['Heading 2'].paragraph_format.space_before = Pt(10)
styles['Heading 2'].paragraph_format.space_after = Pt(5)
for name in ['Code', 'Console']:
    s = styles.add_style(name, 1)
    s.font.name = 'Consolas'
    s.font.size = Pt(8.5 if name == 'Code' else 9)
    s.paragraph_format.line_spacing = Pt(10.5 if name == 'Code' else 11)
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after = Pt(0)
    s.paragraph_format.widow_control = False

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.style = styles['Normal']
run = footer.add_run()
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)

def p(text='', style=None):
    return doc.add_paragraph(text, style)

def h(text, level=1, new=False):
    heading = doc.add_heading(text, level)
    if new:
        heading.paragraph_format.page_break_before = True

def code(text, style='Code', keep_all=False):
    lines = text.rstrip().splitlines()
    for i, line in enumerate(lines):
        para = p(line.expandtabs(4) or ' ', style)
        para.paragraph_format.keep_with_next = (
            (keep_all and i < len(lines) - 1) or
            i < min(5, len(lines) - 1) or
            (i >= len(lines) - 3 and i < len(lines) - 1)
        )
        if i == len(lines) - 1:
            para.paragraph_format.space_after = Pt(6)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    pr.append(borders)
    for j, text in enumerate(headers):
        t.rows[0].cells[j].text = text
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = str(val)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, row in enumerate(t.rows):
        cant = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant)
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), '85')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E7EDF3' if i == 0 else ('F7F8FA' if i % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.02
                for r in para.runs:
                    r.font.size = Pt(9)
                    r.bold = i == 0
    p('').paragraph_format.space_after = Pt(0)
    return t

identity_file = ROOT / 'work' / 'identity.json'
identity = json.loads(identity_file.read_text(encoding='utf-8')) if identity_file.exists() else {}
student = identity.get('name', '________________________________')
group = identity.get('group', '________________________________')
p('Assignment 3', 'Subtitle')
p('Implementing Interfaces and Programming with Mutable and Immutable Models', 'Title')
p('Academic week 3', 'Subtitle')
p(f'Student: {student}\nAcademic group: {group}')
if identity.get('course'):
    p('Course: ' + identity['course'])
if identity.get('instructor'):
    p('Instructor: ' + identity['instructor'])

h('Aim and result')
p('This work redesigns a bookstore inventory model so that shared references cannot silently change stock. The solution uses an immutable C# record, ISBN-based equality, price-based ordering, and validated stock operations that return new records. It reproduces the mutable baseline, contrasts default and custom equality, and demonstrates why filtering by value equality removes more than one duplicate entry.')
p('The supplied catalogue is retained in its original order and with its original data. A separate sorted list is used for display and a further copy is used for the duplicate experiment. The application and checks target .NET 10 and use only the standard library.')
h('Required catalogue', 2)
table(['Position', 'Title', 'ISBN', 'Price', 'Stock'], [
    ['1', 'Refactoring', '111', '45.00', '4'],
    ['2', 'Clean Code', '222', '35.50', '2'],
    ['3', 'The Pragmatic Programmer', '333', '40.00', '6'],
], [.55, 3.1, .7, .7, .72])
h('Submission contents', 2)
p('The source files, project files, verification script and captured console logs accompany this report. Appendix A contains complete text listings of the application, baseline and checks. The report covers the three tasks, expected and actual results, reproducible commands, and questions for the mandatory oral defence.')

h('Task 1 Immutable records and equality', new=True)
h('Why the shared reference prints zero', 2)
p('The statement Book display = featured; copies a reference, not the Book object. In the starter, catalog[0], featured and display therefore refer to the same instance. The exact line display.StockCount = 0; writes to that instance through a public setter. Console.WriteLine(featured.StockCount); reads the same property and prints 0. Mutable means that the object’s observable state can be changed after construction.')
p('The starter code is retained in baseline/OriginalSnippet.txt with its comment removed. Its pasted ordering places a type declaration before top-level statements, which causes compiler error CS8803 in a standard console project. The runnable baseline keeps the same mutable class and executable statements inside Program.Main. This packaging adjustment does not change the aliasing behaviour. The actual baseline output is 0. [8]')
h('What makes the redesigned model immutable', 2)
p('Book is a sealed record class, so it remains a reference type. Merely replacing class with record while retaining public setters would still allow the same bug. Here Title, Isbn and Price have getter-only properties; StockCount has a private init accessor. A validated constructor creates the initial state, and only the record’s own checked methods can initialize a different stock value in a with expression. Strings are immutable and the remaining fields are value types, so the shallow record copy contains no mutable child object. [2, 10]')
p('When display = featured is followed by display = display.Sell(4), display is rebound to a new record. The original featured record and catalog[0] keep stock 4; display has stock 0. Shared references are harmless because the old object cannot be edited through the public API. The class is sealed to keep the equality contract and construction rules closed to derived record types.')
h('Default equality prediction and measured comparison', 2)
p('Prediction: for DefaultBook, the compiler-generated equality compares the runtime record type and every instance field. Its auto-properties store Title, Isbn, Price and StockCount in backing fields, so all four values participate. Two default records with ISBN 111 but different price or stock should be unequal. The comparison record in DefaultBook.cs checks that prediction by execution. [1]')
p('Book implements the record’s IEquatable<Book> contract by defining Equals(Book?) using StringComparer.Ordinal on Isbn only. GetHashCode uses that same comparer and ISBN, which keeps equal records consistent in HashSet<Book>. The compiler supplies object equality and equality operators using this custom typed equality. Same ISBN means equal even if title, price and stock differ; another ISBN means unequal. Matching is exact and case-sensitive, with no trimming or ISBN normalization. [1, 3]')
table(['Comparison', 'Expected', 'Actual'], [
    ['Default records with same ISBN but different price or stock', 'False', 'False'],
    ['Custom records with same ISBN but different other data', 'True', 'True'],
    ['Custom records with different ISBN values', 'False', 'False'],
], [4.85, .66, 1.06])
p('Complete implementations: Appendix A, Book.cs and DefaultBook.cs. The supplied baseline and executable wrapper are also listed there.')

h('Task 2 Ordering and controlled changes', new=True)
h('Ordering without changing the catalogue', 2)
p('Book implements IComparable<Book>. CompareTo returns the result of Price.CompareTo(other.Price), with any non-null book ordered after null. Price comparison is ascending; a negative result means cheaper, zero means the same price, and a positive result means more expensive. No ISBN tie-breaker is added. [4]')
p('The application creates catalog.OrderBy(b => b).ToList(). The default comparer invokes IComparable<Book>, so this actually exercises the implemented interface. OrderBy is stable for equal keys, and ToList materializes a separate list. Anyone holding the original List<Book> sees its unchanged order 111, 222, 333. Callers using the sorted list see 222, 333, 111. Both lists initially reference the same immutable records, which is safe. Sorting catalog itself would change the order visible through every alias of that list. [5, 11]')
p('Equality and ordering deliberately use different attributes because the assignment requires them: two equal ISBNs can compare as different prices, while different ISBNs with the same price compare as zero. Price-based CompareTo is therefore unsuitable for ISBN uniqueness in SortedSet<Book>; use HashSet<Book> with the ISBN equality rule instead. The price comparer remains suitable for catalogue display. [12]')
h('Checked transitions and the stock invariant', 2)
p('The invariant is 0 <= StockCount <= int.MaxValue. Construction rejects negative stock. Restock and Sell reject quantities less than or equal to zero. Sell also rejects a quantity greater than the current stock. Restock evaluates checked(StockCount + quantity), which throws OverflowException if the sum exceeds Int32 capacity instead of assigning a wrapped negative value. Each successful method returns this with { StockCount = ... }; the original instance remains unchanged. [9]')
table(['Operation', 'Validation', 'Failure type'], [
    ['Construction', 'Stock cannot be negative; title and ISBN must be nonblank; price cannot be negative', 'ArgumentException or ArgumentOutOfRangeException'],
    ['Restock', 'Quantity must be positive; checked addition must fit in Int32', 'ArgumentOutOfRangeException or OverflowException'],
    ['Sell', 'Quantity must be positive', 'ArgumentOutOfRangeException'],
    ['Sell', 'Quantity cannot exceed stock', 'InvalidOperationException'],
], [1.02, 3.66, 1.89])
h('Why callers must capture the result', 2)
p('Returning a new record preserves the previous state for other consumers and makes a change explicit at the call site. A caller must use book = book.Sell(quantity), store the result elsewhere, or replace the appropriate list element. Calling book.Sell(1) and discarding its result leaves book unchanged, although invalid requests still throw. This design does not automatically update a list or a database.')
p('A public init accessor would prohibit ordinary later assignments but allow external code such as book with { StockCount = -1 }. That expression would bypass Restock and Sell, and the ordinary validating constructor would not revalidate the altered stock. private init closes this route. Compile-failure checks exercise direct assignment, an external with initializer, and an object initializer. [2, 10]')
p('Complete implementation: Appendix A, Book.cs. The demo and automated checks show new references and unchanged source records.')

h('Task 3 Catalogue report and duplicate handling', new=True)
h('Sorted display and the duplicate experiment', 2)
table(['Display order', 'Title', 'Price', 'Stock'], [
    ['222', 'Clean Code', '35.50', '2'],
    ['333', 'The Pragmatic Programmer', '40.00', '6'],
    ['111', 'Refactoring', '45.00', '4'],
], [.92, 3.5, 1.05, .7])
p('The sorted list is copied before adding a separate Book with ISBN 111 and different price and stock. The existing ISBN 111 record and the added record are different objects but compare equal. Equality-based duplicate detection finds the repeated ISBN. Only the experiment list grows from three to four entries.')
code('var naive = withDuplicate.Where(b => !b.Equals(duplicate)).ToList();')
p('Equals compares ISBN, so the predicate is false for both ISBN 111 records. The naive result contains only Clean Code (222) and The Pragmatic Programmer (333). Its count is 2, not 3. This filter removes an entire equality class; it does not mean “remove the object I just appended”. [6]')
code('var corrected = withDuplicate\n    .Where(b => !ReferenceEquals(b, duplicate)).ToList();')
p('ReferenceEquals tests object identity. Only the exact appended duplicate object fails this predicate, leaving three records. The original Refactoring object survives with price 45.00 and stock 4. The count falls by exactly one, and the original catalogue remains 111, 222, 333 with all supplied values. List<T>.Remove is not used. If the exact same object reference occurred twice, this predicate would exclude both occurrences; the demonstrated duplicate is a separately constructed object. [7]')
table(['State or filter', 'Count', 'Surviving ISBNs'], [
    ['Sorted catalogue', '3', '222, 333, 111'],
    ['Experiment copy with duplicate', '4', '222, 333, 111, 111'],
    ['Where with negated Equals', '2', '222, 333'],
    ['Where with negated ReferenceEquals', '3', '222, 333, 111'],
], [3.4, .65, 2.52])
h('Stock change demonstration', 2)
p('Starting from Refactoring with stock 4, Restock(3) returns stock 7; Sell(2) on that result returns stock 5. Both earlier records retain their values. Sell(6) is then rejected with “Cannot sell 6 copies; only 5 are in stock.” The last valid record keeps stock 5, while the catalogue record remains at 4. The complete, actual console transcript appears in the results appendix.')
p('Complete implementation: Appendix A, Program.cs. Boundary cases, collection preservation and duplicate-filter results are checked in Checks.cs.')

h('Verification and execution', new=True)
h('Build and run steps', 2)
p('Install a .NET 10 SDK, then open PowerShell in the extracted source directory, which contains IFP3.csproj. The tested SDK was 10.0.401 on Windows. No NuGet packages or external services are required. Run these commands in order:')
code('dotnet --version\ndotnet build .\\IFP3.csproj -c Release\ndotnet run --project .\\baseline\\Baseline.csproj -c Release\ndotnet run --project .\\IFP3.csproj -c Release --no-build\ndotnet run --project .\\IFP3.csproj -c Release --no-build -- --check\npowershell -NoProfile -ExecutionPolicy Bypass -File .\\verify.ps1', 'Console')
p('The final command runs the verification workflow, including deliberate compilation failures for forbidden setters. These probe failures are expected evidence that the public API prevents mutation; the script itself must finish successfully. A successful build and a successful run both return exit code 0. The baseline prints 0. The application prints the task demonstrations, and the check runner reports passed checks.')
h('Expected and actual results', 2)
table(['Input or scenario', 'Expected', 'Actual result'], [
    ['Shared mutable alias', 'featured stock becomes 0', '0'],
    ['Immutable alias with captured Sell(4)', 'Source 4 and new value 0', '4 and 0'],
    ['Original and sorted catalogue ISBN order', '111,222,333 / 222,333,111', 'Matches'],
    ['Duplicate exclusion from four entries', 'Equals gives 2; identity gives 3', '2 and 3'],
    ['Initial stock -1', 'Reject construction', 'ArgumentOutOfRangeException'],
    ['Zero or negative restock and sale', 'Reject quantity', 'ArgumentOutOfRangeException'],
    ['Sell all available stock', 'Return stock 0; preserve source', 'Passed'],
    ['Oversell by one', 'Reject; keep prior stock', 'InvalidOperationException'],
    ['Restock beyond int.MaxValue', 'Reject the overflowing sum', 'OverflowException'],
    ['Discard a valid Sell result', 'Original stock unchanged', 'Passed'],
    ['External stock assignment and initializers', 'Compilation rejected', 'Rejected'],
], [2.5, 2.1, 1.97])
p('The Release build completed with 0 warnings and 0 errors. All 43 runtime checks and all 6 compile-time guards passed. Actual evidence is captured in evidence/build.txt, evidence/baseline.txt, evidence/demo.txt, evidence/checks.txt and evidence/compile-guards.txt. Appendix B reproduces the console demonstration and check output. The verification also checks equality operators, hashing, null handling, price ties, valid integer boundaries and unchanged catalogue data.')

h('Architecture and discussion', new=True)
h('Structure and complexity', 2)
p('Book owns the data rules. Program composes the catalogue operations and prints evidence. DefaultBook provides an isolated comparison with generated equality. Checks supplies dependency-free automated assertions, while verify.ps1 exercises build, execution and compile-time access restrictions. The mutable baseline is a separate project so the two Book models cannot conflict.')
p('Let n be the catalogue size and L the ISBN length. Stock operations and price comparisons take O(1) time and allocate at most one fixed-size record, O(1) extra space. ISBN equality is O(L) in the worst case; hashing the ISBN is O(L). Sorting and materializing the view takes O(n log n) time and O(n) auxiliary space. Counting matches or filtering with Equals takes O(nL) time; reference filtering takes O(n). Counting uses O(1) auxiliary space, whereas ToList needs O(n) result storage. For fixed-size ISBNs, L can be treated as a constant. Printing takes time proportional to the amount of text produced.')
h('Questions for oral defence', 2)
for question, answer in [
    ('Why does the starter print 0?', 'display and featured reference the same mutable object. The public setter at display.StockCount = 0 changes that one object.'),
    ('Does record mean immutable?', 'No. A record class is still a reference type and can contain setters or mutable referenced children. Restricted accessors and immutable fields provide immutability here.'),
    ('Why customize both equality and hashing?', 'The identity rule is ISBN only. Equal ISBNs must produce equal hashes so hash-based collections use the same identity rule.'),
    ('Why can CompareTo return zero for unequal books?', 'Ordering uses price while equality uses ISBN. This is appropriate for the required display, but price-based ordered sets do not enforce ISBN identity.'),
    ('What happens when a Sell result is ignored?', 'A successful new record is discarded; the old variable and any list entry retain their original stock. A failing sale still throws.'),
    ('Why is public init unsafe for stock?', 'An external with initializer could set negative stock without the validation in Sell or Restock. The normal constructor is not re-run for that change.'),
    ('Why does the first duplicate filter remove two books?', 'Both objects have the same ISBN, so both satisfy Equals(duplicate). ReferenceEquals selects only the separate appended object.'),
    ('What small change could an instructor request?', 'To display highest price first, change OrderBy(b => b) to OrderByDescending(b => b). Keep CompareTo ascending and keep ToList so the original list order remains unchanged.'),
]:
    q = p()
    q.add_run(question + ' ').bold = True
    q.add_run(answer)
h('Conclusion', 2)
p('The redesign preserves the required dataset, prevents hidden changes through aliases, and makes stock changes explicit and validated. The experiments distinguish record equality, ISBN identity, price ordering and object identity. These distinctions explain both the original mutation bug and the duplicate-filtering trap.')

h('Appendix A Complete source code', new=True)
p('These listings reproduce the submitted files as readable text. Book.cs implements Tasks 1 and 2; Program.cs demonstrates all tasks. The baseline, comparison record, checks, verification script and project files make the submission reproducible.')
files = [
    ('Book model for Tasks 1 and 2', 'Book.cs'),
    ('Default equality comparison for Task 1', 'DefaultBook.cs'),
    ('Application for Tasks 1 to 3', 'Program.cs'),
    ('Automated checks', 'Checks.cs'),
    ('Verification script', 'verify.ps1'),
    ('Main project configuration', 'IFP3.csproj'),
    ('Supplied starter', 'baseline/OriginalSnippet.txt'),
    ('Runnable baseline wrapper', 'baseline/Program.cs'),
    ('Baseline project configuration', 'baseline/Baseline.csproj'),
]
for label, path in files:
    h(label, 2)
    pp = p(path)
    pp.paragraph_format.keep_with_next = True
    pp.paragraph_format.space_after = Pt(4)
    code(read(ROOT / path), keep_all=(path == 'baseline/OriginalSnippet.txt'))

h('Appendix B Actual execution evidence', new=True)
for title, path in [
    ('Mutable baseline output', 'evidence/baseline.txt'),
    ('Application output', 'evidence/demo.txt'),
    ('Automated check output', 'evidence/checks.txt'),
]:
    h(title, 2)
    text = read(ROOT / path).strip()
    code(text, 'Console')

h('Compile time guard results', 2)
p('The valid control caller compiled successfully. The invalid caller failed with exit code 1 and six errors. The following diagnostic codes are extracted from the actual compiler log; full messages and paths are retained in evidence/compile-guards.txt.')
guard_text = read(ROOT / 'evidence/compile-guards.txt')
guard_codes = {int(line): error for line, error in re.findall(r'Probe\.cs\((\d+),\d+\): error (CS\d+)', guard_text)}
table(['Probe line', 'Forbidden operation', 'Actual diagnostic'], [
    [str(line), label, guard_codes.get(line, 'MISSING')]
    for line, label in [
        (3, 'Assign StockCount after construction'),
        (4, 'Set StockCount in an external with expression'),
        (5, 'Set StockCount in an object initializer'),
        (6, 'Change Title in an external with expression'),
        (7, 'Change Isbn in an external with expression'),
        (8, 'Change Price in an external with expression'),
    ]
], [.8, 4.32, 1.45])
code('All 6 compile guards passed (expected errors).', 'Console')

h('References', new=True)
refs = [
    ('1', 'Microsoft Learn', 'Records in C#', 'https://learn.microsoft.com/en-us/dotnet/csharp/language-reference/builtin-types/record'),
    ('2', 'Microsoft Learn', 'The with expression', 'https://learn.microsoft.com/en-us/dotnet/csharp/language-reference/operators/with-expression'),
    ('3', 'Microsoft Learn', 'Object.GetHashCode method', 'https://learn.microsoft.com/en-us/dotnet/api/system.object.gethashcode'),
    ('4', 'Microsoft Learn', 'IComparable<T>.CompareTo method', 'https://learn.microsoft.com/en-us/dotnet/api/system.icomparable-1.compareto'),
    ('5', 'Microsoft Learn', 'Enumerable.OrderBy method', 'https://learn.microsoft.com/en-us/dotnet/api/system.linq.enumerable.orderby'),
    ('6', 'Microsoft Learn', 'Enumerable.Where method', 'https://learn.microsoft.com/en-us/dotnet/api/system.linq.enumerable.where'),
    ('7', 'Microsoft Learn', 'Object.ReferenceEquals method', 'https://learn.microsoft.com/en-us/dotnet/api/system.object.referenceequals'),
    ('8', 'Microsoft Learn', 'Resolve errors related to the Main method and top level statements', 'https://learn.microsoft.com/en-us/dotnet/csharp/language-reference/compiler-messages/entry-point-errors'),
    ('9', 'Microsoft Learn', 'Checked and unchecked statements', 'https://learn.microsoft.com/en-us/dotnet/csharp/language-reference/statements/checked-and-unchecked'),
    ('10', 'Microsoft Learn', 'The init keyword', 'https://learn.microsoft.com/en-us/dotnet/csharp/language-reference/keywords/init'),
    ('11', 'Microsoft Learn', 'Comparer<T>.Default property', 'https://learn.microsoft.com/en-us/dotnet/api/system.collections.generic.comparer-1.default'),
    ('12', 'Microsoft Learn', 'IEquatable<T> interface', 'https://learn.microsoft.com/en-us/dotnet/api/system.iequatable-1'),
]
for number, publisher, title, url in refs:
    p(f'[{number}] {publisher}. {title}.')
    q = p(url)
    for run in q.runs:
        run.font.size = Pt(9)
p('Assignment source: Assignment 3 Week 3, Implementing Interfaces and Programming with a Mutable Type/Immutable Models, supplied task sheet.')

doc.core_properties.title = 'Assignment 3 Implementing Interfaces and Immutable Book Models'
doc.core_properties.subject = 'C# records, interfaces, equality, ordering and checked stock transitions'
doc.core_properties.author = identity.get('name', '')
doc.core_properties.keywords = 'C#, records, ISBN, immutability, interfaces'
doc.core_properties.comments = ''
doc.core_properties.created = datetime(2026, 9, 27, tzinfo=timezone.utc)
doc.core_properties.modified = datetime(2026, 9, 27, tzinfo=timezone.utc)
for root_element in [doc.styles.element, doc._element]:
    for border in list(root_element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
